"""Gera o relatório de entrega do Trabalho 2 em Word."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Relatorio_Trabalho_2_Nexus_Heroes.docx"
DOC = Document()
SEC = DOC.sections[0]
SEC.page_width = Cm(21)
SEC.page_height = Cm(29.7)
SEC.top_margin = Cm(2.2)
SEC.bottom_margin = Cm(2.0)
SEC.left_margin = Cm(3.0)
SEC.right_margin = Cm(2.4)
SEC.header_distance = Cm(1.0)
SEC.footer_distance = Cm(1.0)
SEC.different_first_page_header_footer = True


def set_font(style, size, bold=False):
    style.font.name = "Aptos"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Aptos")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Aptos")


normal = DOC.styles["Normal"]
set_font(normal, 10.5)
normal.paragraph_format.line_spacing = 1.17
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True

title = DOC.styles["Title"]
set_font(title, 18, True)
title.paragraph_format.space_after = Pt(12)
title.paragraph_format.keep_with_next = True
title_ppr = title._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)

for name, size, before, after in [
    ("Heading 1", 13, 14, 7),
    ("Heading 2", 11, 10, 5),
]:
    style = DOC.styles[name]
    set_font(style, size, True)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True


def add_par(text="", *, bold_start=None, align=None, style=None, keep=False):
    p = DOC.add_paragraph(style=style)
    if bold_start and text.startswith(bold_start):
        p.add_run(bold_start).bold = True
        p.add_run(text[len(bold_start):])
    else:
        p.add_run(text)
    if align is not None:
        p.alignment = align
    elif style not in ("Title", "Heading 1", "Heading 2"):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = keep
    return p


def heading(text, level=1):
    return add_par(text, style=f"Heading {level}")


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:color"), color)
        el.set(qn("w:sz"), size)
        tc_borders.append(el)


def margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    m = tc_pr.first_child_found_in("w:tcMar")
    if m is None:
        m = OxmlElement("w:tcMar")
        tc_pr.append(m)
    for name, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        m.append(el)


def table(headers, rows, widths):
    t = DOC.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Cm(width)
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.width = Cm(widths[i])
        cell.text = h
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "213B56")
        margins(cell)
        borders(cell)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9.5)
    for n, data in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(data):
            c = cells[i]
            c.width = Cm(widths[i])
            c.text = str(value)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n % 2:
                shade(c, "F3F6F9")
            margins(c)
            borders(c)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(9.3)
    for row in t.rows:
        row._tr.get_or_add_trPr()
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    DOC.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def placeholder(caption, height_cm):
    p = DOC.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(caption).bold = True
    t = DOC.add_table(rows=1, cols=1)
    row = t.rows[0]
    row.height = Cm(height_cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = row.cells[0]
    borders(cell, "AEB8C2", "8")
    margins(cell, 120, 120, 120, 120)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run("Inserir aqui")
    r.font.color.rgb = RGBColor(120, 130, 140)
    r.font.italic = True
    DOC.add_paragraph().paragraph_format.space_after = Pt(0)


def link_line(label):
    p = DOC.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(9)
    p.add_run(label + " ").bold = True
    p.add_run("____________________________________________________________")


def page_break():
    DOC.add_page_break()


# Capa
add_par("[NOME DA INSTITUIÇÃO]", align=WD_ALIGN_PARAGRAPH.CENTER)
add_par("[CURSO]", align=WD_ALIGN_PARAGRAPH.CENTER)
add_par("Programação Orientada a Objetos", align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(5):
    DOC.add_paragraph()
add_par("TRABALHO 2 PROGRAMAÇÃO ORIENTADA A OBJETOS", align=WD_ALIGN_PARAGRAPH.CENTER, style="Title")
add_par("Análise e modelagem do jogo Nexus Heroes", align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(4):
    DOC.add_paragraph()
add_par("Eric Dias", align=WD_ALIGN_PARAGRAPH.CENTER)
add_par("Matheus Howe Habeck", align=WD_ALIGN_PARAGRAPH.CENTER)
add_par("Professor(a): [NOME]", align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(4):
    DOC.add_paragraph()
add_par("[CIDADE]", align=WD_ALIGN_PARAGRAPH.CENTER)
add_par("[DATA]", align=WD_ALIGN_PARAGRAPH.CENTER)
page_break()

# Introdução e etapa 1
heading("1 Introdução")
add_par("Este trabalho apresenta a análise do domínio de Nexus Heroes, a modelagem de classes UML e sua implementação em Java. A proposta é demonstrar abstração, encapsulamento, herança, sobrescrita de métodos e preservação das regras de estado observadas no jogo. O ambiente foi descrito como um labirinto 3D isométrico com interações de combate, coleta e progressão; a implementação trata dessas regras no console, sem reproduzir a interface visual.")

heading("2 Exploração do ambiente")
heading("2.1 Ambiente e objetivo", 2)
add_par("O personagem percorre um labirinto, encontra inimigos e elementos do mapa e busca alcançar o portal que conclui a fase. Inimigos vivos bloqueiam o avanço, podem contra-atacar e concedem Hero Coins quando derrotados. Baús concedem experiência e moedas; itens recuperam recursos; armadilhas retiram vida.")
add_par("A descrição disponível não apresenta um mapa completo, dimensões, coordenadas ou caminhos específicos. Por isso, esta análise registra as interações informadas sem atribuir percursos ou cenários não documentados ao jogo.")

heading("2.2 Personagens e atributos", 2)
table(["Classe", "HP máximo", "Mana máxima", "ATK espada", "MATK magia"], [
    ["Guerreiro", "120", "40", "25", "10"],
    ["Mago", "80", "120", "15", "40"],
], [3.2, 2.6, 3.2, 3.0, 3.0])
add_par("Os estados identificados são nome, vida, vida máxima, mana, mana máxima, ataque, nível, experiência (XP), moedas e posição. Goblin e Golem aparecem como inimigos, mas seus valores de vida, ataque e recompensa não foram informados.")

heading("2.3 Entidades e comportamentos", 2)
table(["Entidade", "Papel no domínio", "Ações relacionadas"], [
    ["Personagem", "Herói que combate e progride", "Andar, atacar, coletar, receber dano, subir de nível"],
    ["Inimigo", "Oponente que bloqueia o caminho", "Calcular dano, contra-atacar, entregar moedas"],
    ["Item e HeroCoin", "Recursos coletáveis", "Aplicar bônus ou acrescentar moedas"],
    ["Armadilha", "Evento de dano", "Retirar 20 HP e sinalizar o evento"],
    ["Baú", "Fonte de recompensa", "Conceder XP e Hero Coins"],
    ["Portal e Mapa", "Objetivo e navegação", "Controlar passagem e conclusão da fase"],
], [3.1, 5.2, 7.0])

page_break()
heading("2.4 Interações e regras de transição", 2)
table(["Elemento", "Efeito informado", "Regra de estado"], [
    ["Cristal de Mana", "+25 mana", "Não ultrapassa a mana máxima"],
    ["Orbe de Vida", "+20 HP", "Não ultrapassa a vida máxima"],
    ["Inimigo", "Combate e Hero Coins", "Bloqueia enquanto vivo; recompensa após derrota"],
    ["Armadilha", "-20 HP", "Vida não fica abaixo de zero"],
    ["Baú", "+XP e +moedas", "Recompensa associada à abertura"],
    ["Portal", "Vitória", "Entrada conclui a fase"],
], [3.3, 3.6, 8.4])
add_par("Em todos os estados, 0 ≤ vida ≤ vida máxima e 0 ≤ mana ≤ mana máxima. Ao subir de nível, a vida e a mana são restauradas e o ataque aumenta. Os valores de XP por nível, incremento de ataque e recompensas não constam da descrição e foram mantidos configuráveis na implementação.")

placeholder("Foto da exploração do jogo", 4.5)
page_break()

# Modelagem UML
heading("3 Modelagem de classes UML")
add_par("A arquitetura separa combate, efeitos de itens e navegação. Combatente concentra nome, vida e posição; Personagem acrescenta mana, experiência, moedas e ataques; Inimigo define bloqueio, contra-ataque e recompensa. Guerreiro e Mago são subclasses de Personagem, enquanto Goblin e Golem são subclasses de Inimigo. A relação é-um justifica cada generalização.")
add_par("Item representa os efeitos coletáveis, com HeroCoin, CristalMana e OrbeVida como especializações. ElementoMapa estabelece o contrato comum para os elementos encontrados no mapa. Mapa valida o deslocamento e a entrada no portal. Baú contém sua recompensa; Armadilha aplica dano e emite TrapDamageException.")
heading("3.1 Encapsulamento e relações", 2)
add_par("Atributos mutáveis permanecem privados. Construtores exigem dados obrigatórios e validam estados iniciais. Os mutadores de vida e mana preservam seus limites. O diagrama apresenta visibilidades, assinaturas, multiplicidades, notas de invariantes, generalização, realização de interface, associação, agregação, composição e dependências. Os métodos calcularDano, aumentarAtaque e aplicarEm evidenciam a sobrescrita nas subclasses.")
placeholder("Diagrama de classes UML", 7.0)
link_line("Link editável do diagrama UML:")

# Implementação
page_break()
heading("4 Implementação em Java")
add_par("O projeto foi implementado com Java 21 como alvo de compilação e Maven para organização, compilação e testes. As subclasses usam extends e chamam super nos construtores. Todos os métodos sobrescritos recebem @Override. A classe Main demonstra no console a criação de objetos, o ataque, o contra-ataque, a coleta, a progressão de nível e a conclusão da fase.")
add_par("O ataque do Guerreiro usa a espada; o ataque principal do Mago usa MATK. O inimigo sobrevivente contra-ataca na simulação. A derrota entrega HeroCoin uma única vez. O baú concede XP e moedas; a subida de nível restaura HP e mana e aumenta o ataque principal. Cristal e orbe respeitam os máximos. A armadilha aplica 20 de dano e sinaliza TrapDamageException.")
heading("4.1 Validação", 2)
add_par("Foram executados 11 testes unitários com JUnit 5, sem falhas nem erros. Eles verificam construção válida, rejeição de parâmetros inválidos, limites de vida e mana, transições permitidas e proibidas, combate, contra-ataque, recompensa única, baú, armadilha, portal e polimorfismo. A simulação de console também foi executada.")
heading("4.2 Repositório do projeto", 2)
add_par("O repositório deve conter o código-fonte Java, os testes, a configuração Maven e o diagrama de classes editável.")
link_line("Link do GitHub do projeto:")

heading("5 Considerações finais")
add_par("O modelo implementa as regras de domínio informadas para Nexus Heroes e mantém válidos os limites de vida e mana em suas transições. A hierarquia de combatentes permite comportamento polimórfico sem duplicar o estado comum. A exploração integral de caminhos não pôde ser afirmada porque o mapa completo e o executável do jogo não foram disponibilizados; os números ausentes permanecem parâmetros explícitos, sem serem tratados como regras oficiais.")

heading("Referências")
add_par("Enunciado da atividade Trabalho 2 Programação Orientada a Objetos, fornecido pela disciplina.")
add_par("Descrição de Nexus Heroes e dados da Etapa 1 fornecidos pelos integrantes.")


# Número de página discreto no rodapé.
footer = SEC.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Página ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

DOC.core_properties.title = "Trabalho 2 Programação Orientada a Objetos Nexus Heroes"
DOC.core_properties.subject = "Análise modelagem UML e implementação Java"
DOC.core_properties.author = "Eric Dias e Matheus Howe Habeck"
DOC.save(OUT)
print(OUT)
